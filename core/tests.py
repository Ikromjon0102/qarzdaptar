import json
from unittest import mock

from django.contrib.auth.models import User
from django.test import TestCase, TransactionTestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from .models import BotSignup, Client, Debt, Shop, UserProfile
from .telegram_auth import sign_init_data, verify_init_data

TEST_BOT_TOKEN = '123456:TEST-token'


def signed_init(telegram_id, auth_date=None, token=TEST_BOT_TOKEN):
    """Telegram imzolagandek initData (test uchun)."""
    import time
    return sign_init_data({
        'auth_date': str(int(auth_date or time.time())),
        'query_id': 'AAE',
        'user': json.dumps({'id': telegram_id, 'first_name': 'Test'}),
    }, bot_token=token)


def make_shop(tg_id='111111', name="Test Do'kon"):
    user = User.objects.create_user(username=tg_id, password='1')
    shop = Shop.objects.create(name=name, owner=user)
    UserProfile.objects.create(user=user, shop=shop, role='admin')
    return user, shop


@mock.patch('core.bot_signup.telegram')
class BotSignupTests(TestCase):
    """Botda do'kon ochish: /start signup -> nomi -> turi -> do'kon tayyor."""

    def update(self, payload):
        return self.client.post(reverse('telegram_webhook'), data=json.dumps(payload),
                                content_type='application/json')

    def text(self, text, chat_id=555):
        return self.update({'message': {'chat': {'id': chat_id}, 'from': {'id': chat_id, 'first_name': 'Ali'},
                                        'text': text}})

    def tap(self, data, chat_id=555):
        return self.update({'callback_query': {'id': 'cb', 'data': data, 'from': {'id': chat_id, 'first_name': 'Ali'},
                                               'message': {'chat': {'id': chat_id}, 'message_id': 9}}})

    def test_full_flow_creates_shop_with_category_and_trial(self, tg):
        self.text('/start signup')
        self.assertIn('1/2', tg.send_message.call_args[0][1])
        self.text('  Baraka   Market ')
        self.assertIn('2/2', tg.send_message.call_args[0][1])
        buttons = [b['callback_data'] for row in tg.send_message.call_args[1]['reply_markup']['inline_keyboard']
                   for b in row]
        self.assertIn('signup_cat:plumbing', buttons)
        with mock.patch('core.views.answer_callback'):
            self.tap('signup_cat:plumbing')

        shop = Shop.objects.get(owner__username='555')
        self.assertEqual(shop.name, 'Baraka Market')
        self.assertEqual(shop.category, 'plumbing')
        self.assertGreaterEqual(shop.days_left, 13)
        self.assertEqual(shop.owner.profile.role, 'admin')
        self.assertTrue(shop.settings.usd_rate)
        self.assertIn('ochildi', tg.edit_message.call_args[0][2])
        self.assertIn('web_app', str(tg.edit_message.call_args[1]['reply_markup']))
        self.assertFalse(BotSignup.objects.exists())

    def test_invalid_name_is_asked_again(self, tg):
        self.text('/start signup')
        self.text('A')
        self.assertEqual(BotSignup.objects.get().step, 'name')

    def test_registered_user_gets_menu_instead(self, tg):
        make_shop(tg_id='555')
        with mock.patch('core.views.send_menu') as menu:
            self.text('/start signup')
        menu.assert_called_once()
        self.assertFalse(BotSignup.objects.exists())

    def test_stale_category_tap_does_nothing(self, tg):
        with mock.patch('core.views.answer_callback'):
            self.tap('signup_cat:grocery')
        self.assertFalse(Shop.objects.exists())

    def test_cancel(self, tg):
        self.text('/start signup')
        self.text('/cancel')
        self.assertFalse(BotSignup.objects.exists())
        self.text('Baraka')  # endi oddiy matn hech narsa yaratmaydi
        self.assertFalse(BotSignup.objects.exists())

    def test_plain_start_offers_signup_button(self, tg):
        with mock.patch('core.views.telegram') as views_tg:
            self.text('/start')
        markup = views_tg.send_message.call_args[1]['reply_markup']
        self.assertEqual(markup['inline_keyboard'][0][0]['callback_data'], 'signup_start')


@mock.patch('requests.post')
class SubscriptionBannerTests(TestCase):
    def test_menu_hides_banner_without_subscription_date(self, _post):
        user, _shop = make_shop()
        self.client.force_login(user)
        resp = self.client.get(reverse('main_menu'))
        self.assertNotContains(resp, 'Obuna tugashiga')


@mock.patch('core.views.send_tg_msg')
@mock.patch('core.views.send_menu')
class WebhookTests(TestCase):
    def post_text(self, text, chat_id=555):
        return self.client.post(
            reverse('telegram_webhook'),
            data=json.dumps({'message': {'chat': {'id': chat_id}, 'text': text}}),
            content_type='application/json',
        )

    def test_start_login_sends_menu_not_error(self, send_menu, send_msg):
        self.post_text('/start login')
        send_menu.assert_called_once()
        send_msg.assert_not_called()

    def test_id_command_sends_telegram_id(self, send_menu, send_msg):
        self.post_text('/id', chat_id=777)
        self.assertIn('777', send_msg.call_args[0][1])

    def test_invite_token_links_client(self, send_menu, send_msg):
        _user, shop = make_shop()
        c = Client.objects.create(shop=shop, full_name='Vali', phone='+998901112233')
        self.post_text(f'/start {c.invite_token}', chat_id=999)
        c.refresh_from_db()
        self.assertEqual(c.telegram_id, 999)


@mock.patch('requests.post')
class ClientAndDebtTests(TestCase):
    def setUp(self):
        self.user, self.shop = make_shop()
        self.client.force_login(self.user)

    def sale(self, **extra):
        data = {
            'sale_mode': 'debt', 'payment_type': 'cash',
            'product_name[]': ['Non'], 'quantity[]': ['2'], 'price[]': ['5000'], 'currency[]': ['uzs'],
        }
        data.update(extra)
        return self.client.post(reverse('create_debt'), data)

    def test_same_phone_allowed_in_two_shops(self, _post):
        _u2, shop2 = make_shop(tg_id='222222', name='Ikkinchi')
        Client.objects.create(shop=self.shop, full_name='A', phone='000000000')
        Client.objects.create(shop=shop2, full_name='B', phone='000000000')
        self.assertEqual(Client.objects.filter(phone='000000000').count(), 2)

    def test_cash_sale_works_in_second_shop(self, _post):
        Client.objects.create(shop=self.shop, full_name='Kassa', phone='000000000')
        u2, _shop2 = make_shop(tg_id='222222', name='Ikkinchi')
        self.client.force_login(u2)
        resp = self.sale(sale_mode='cash')
        self.assertEqual(resp.status_code, 302)

    def test_unlinked_client_debt_confirmed_by_choice(self, _post):
        c = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233')
        resp = self.sale(client=c.id, no_tg_action='confirm')
        self.assertRedirects(resp, reverse('admin_client_detail', args=[c.id]))
        self.assertEqual(Debt.objects.get(client=c).status, 'confirmed')

    def test_unlinked_client_debt_can_stay_pending(self, _post):
        c = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233')
        self.sale(client=c.id, no_tg_action='pending')
        self.assertEqual(Debt.objects.get(client=c).status, 'pending')

    def test_linked_client_debt_stays_pending(self, _post):
        c = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233', telegram_id=42)
        self.sale(client=c.id, no_tg_action='confirm')
        self.assertEqual(Debt.objects.get(client=c).status, 'pending')

    def test_status_page_shows_message(self, _post):
        c = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233')
        d = Debt.objects.create(shop=self.shop, client=c, amount_uzs=1000, items='x', status='pending')
        session = self.client.session
        session['client_id'] = c.id  # mijozning o'z sessiyasi
        session.save()
        resp = self.client.post(reverse('debt_detail', args=[d.uuid]), {'action': 'confirm'})
        self.assertContains(resp, 'hisobingizga yozildi')

    def test_confirmed_debt_card_uses_theme_background(self, _post):
        c = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233')
        Debt.objects.create(shop=self.shop, client=c, amount_uzs=1000, items='x', status='confirmed')
        session = self.client.session
        session['client_id'] = c.id
        session.save()
        resp = self.client.get(reverse('client_cabinet'))
        self.assertNotContains(resp, '#232E3C')
        self.assertContains(resp, 'var(--tg-secondary)')


class ParseAmountTests(TestCase):
    def test_formats(self):
        from .utils import parse_amount
        self.assertEqual(parse_amount('1 500 000'), 1500000)
        self.assertEqual(parse_amount('1,500,000'), 1500000)
        self.assertEqual(parse_amount('12,5'), 12.5)
        self.assertEqual(parse_amount('abc'), 0)
        self.assertEqual(parse_amount('-5'), 0)


@mock.patch('requests.post')
class SaleAndPaymentFormTests(TestCase):
    def setUp(self):
        self.user, self.shop = make_shop()
        self.client.force_login(self.user)
        self.vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233')

    def sale(self, **extra):
        data = {
            'sale_mode': 'debt', 'payment_type': 'cash', 'no_tg_action': 'confirm',
            'product_name[]': ['Shakar'], 'quantity[]': ['2'], 'price[]': ['1 500 000'], 'currency[]': ['uzs'],
        }
        data.update(extra)
        return self.client.post(reverse('create_debt'), data)

    def test_formatted_price_is_parsed(self, _post):
        self.sale(client=self.vali.id)
        self.assertEqual(Debt.objects.get(client=self.vali).amount_uzs, 3000000)

    def test_missing_client_keeps_entered_items(self, _post):
        resp = self.sale(client='')
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "mijozni tanlang")
        self.assertContains(resp, 'Shakar')  # prefill JSON ichida
        self.assertFalse(Debt.objects.exists())

    def test_sale_without_items_is_rejected(self, _post):
        resp = self.sale(client=self.vali.id, **{'price[]': ['']})
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(Debt.objects.exists())

    def test_cash_sale_is_separate_from_credit_stats(self, _post):
        from .views import shop_stats
        self.sale(sale_mode='cash', payment_type='transfer')
        self.sale(client=self.vali.id)
        stats = shop_stats(self.shop)
        self.assertEqual(stats['sales_uzs'], 3000000)   # faqat nasiya
        self.assertEqual(stats['income_uzs'], 0)
        self.assertEqual(stats['cash_uzs'], 3000000)
        self.assertEqual(stats['cash_count'], 1)
        self.assertTrue(Debt.objects.filter(is_cash_sale=True, payment_method='transfer').exists())

    def test_cash_client_hidden_from_lists(self, _post):
        self.sale(sale_mode='cash')
        resp = self.client.get(reverse('create_payment'))
        self.assertNotContains(resp, 'Naqd Savdo (Kassa)')
        resp = self.client.get(reverse('dashboard'))
        self.assertNotContains(resp, 'Naqd Savdo (Kassa)')

    def test_payment_with_formatted_amount(self, _post):
        resp = self.client.post(reverse('create_payment'), {
            'client_id': self.vali.id, 'amount_uzs': '250 000', 'payment_method': 'transfer', 'note': 'avans',
        })
        self.assertRedirects(resp, reverse('admin_client_detail', args=[self.vali.id]))
        payment = Debt.objects.get(client=self.vali)
        self.assertEqual(payment.amount_uzs, -250000)
        self.assertEqual(payment.payment_method, 'transfer')
        self.assertIn('avans', payment.items)

    def test_payment_error_keeps_input(self, _post):
        resp = self.client.post(reverse('create_payment'), {'client_id': '', 'amount_uzs': '250 000'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '250 000')
        self.assertFalse(Debt.objects.exists())

    def test_mark_cash_sales_command(self, _post):
        from django.core.management import call_command
        sale = Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=100, items='Non', status='confirmed')
        Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=-100, transaction_type='payment',
                            items=f"To'lov: Non (ID: {sale.id})", status='confirmed')
        call_command('mark_cash_sales', stdout=mock.Mock())
        self.assertEqual(Debt.objects.filter(is_cash_sale=True).count(), 2)


def make_worker(shop, tg_id='333333'):
    from .models import AllowedAdmin
    user = User.objects.create_user(username=tg_id)
    UserProfile.objects.create(user=user, shop=shop, role='worker')
    AllowedAdmin.objects.create(shop=shop, name='Xodim', telegram_id=tg_id)
    return user


@mock.patch('requests.post')
class NavigationAndRoleTests(TestCase):
    def setUp(self):
        self.owner, self.shop = make_shop()
        self.worker = make_worker(self.shop)
        self.vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233')

    def test_worker_blocked_from_admin_sections(self, _post):
        self.client.force_login(self.worker)
        for name in ('settings', 'broadcast', 'admin_control'):
            resp = self.client.get(reverse(name))
            self.assertRedirects(resp, reverse('main_menu'), msg_prefix=name)

    def test_worker_menu_hides_admin_links(self, _post):
        self.client.force_login(self.worker)
        resp = self.client.get(reverse('main_menu'))
        self.assertContains(resp, 'Xodim ·')
        self.assertNotContains(resp, reverse('settings'))
        self.assertNotContains(resp, reverse('admin_control'))
        self.client.force_login(self.owner)
        resp = self.client.get(reverse('main_menu'))
        self.assertContains(resp, reverse('admin_control'))
        self.assertContains(resp, reverse('logout'))

    def test_worker_cannot_delete_transactions(self, _post):
        debt = Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=100, items='x', status='confirmed')
        self.client.force_login(self.worker)
        self.client.get(reverse('manage_debt', args=[debt.uuid, 'delete']))
        self.assertTrue(Debt.objects.filter(id=debt.id).exists())
        self.client.force_login(self.owner)
        self.client.get(reverse('manage_debt', args=[debt.uuid, 'delete']))
        self.assertFalse(Debt.objects.filter(id=debt.id).exists())

    def test_cannot_touch_other_shop_debt(self, _post):
        other_owner, other_shop = make_shop(tg_id='444444', name='Boshqa')
        other_client = Client.objects.create(shop=other_shop, full_name='X', phone='+998900000001')
        debt = Debt.objects.create(shop=other_shop, client=other_client, amount_uzs=1, items='x')
        self.client.force_login(self.owner)
        resp = self.client.get(reverse('manage_debt', args=[debt.uuid, 'delete']))
        self.assertEqual(resp.status_code, 404)

    def test_owner_cannot_be_removed_from_staff(self, _post):
        from .models import AllowedAdmin
        owner_admin = AllowedAdmin.objects.create(shop=self.shop, name='Egasi', telegram_id=int(self.owner.username))
        worker_admin = AllowedAdmin.objects.get(telegram_id=333333)
        self.client.force_login(self.owner)
        self.client.post(reverse('manage_admins_id', args=['delete', owner_admin.id]))
        self.assertTrue(User.objects.filter(id=self.owner.id).exists())
        self.client.get(reverse('manage_admins_id', args=['delete', worker_admin.id]))  # GET bilan o'chmaydi
        self.assertTrue(User.objects.filter(id=self.worker.id).exists())
        self.client.post(reverse('manage_admins_id', args=['delete', worker_admin.id]))
        self.assertFalse(User.objects.filter(id=self.worker.id).exists())

    def test_logout_requires_post(self, _post):
        self.client.force_login(self.owner)
        self.client.get(reverse('logout'))
        self.assertIn('_auth_user_id', self.client.session)
        resp = self.client.post(reverse('logout'))
        self.assertRedirects(resp, reverse('landing_page'))
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_back_button_uses_real_url(self, _post):
        self.client.force_login(self.owner)
        resp = self.client.get(reverse('create_debt'))
        self.assertContains(resp, f'href="{reverse("main_menu")}"')
        self.assertNotContains(resp, 'history.back()')
        resp = self.client.get(reverse('client_edit', args=[self.vali.id]))
        self.assertContains(resp, f'href="{reverse("admin_client_detail", args=[self.vali.id])}"')

    def test_client_list_and_new_client_flow(self, _post):
        Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=500000, items='x', status='confirmed')
        Client.objects.create(shop=self.shop, full_name='Naqd Savdo (Kassa)', phone='000000000')
        self.client.force_login(self.worker)
        resp = self.client.get(reverse('client_list'))
        self.assertContains(resp, '500 000')
        self.assertNotContains(resp, 'Kassa')

        resp = self.client.post(reverse('client_add'), {'full_name': 'Ali', 'phone': '90 123 45 67'})
        ali = Client.objects.get(full_name='Ali')
        self.assertEqual(ali.phone, '+998901234567')
        self.assertRedirects(resp, reverse('admin_client_detail', args=[ali.id]))
        resp = self.client.get(reverse('admin_client_detail', args=[ali.id]))
        self.assertContains(resp, str(Client.objects.get(id=ali.id).invite_token))

    def test_dashboard_shows_top_debtors_only(self, _post):
        Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=500000, items='x', status='confirmed')
        Client.objects.create(shop=self.shop, full_name='Toza Mijoz', phone='+998900000002')
        self.client.force_login(self.owner)
        resp = self.client.get(reverse('dashboard'))
        self.assertContains(resp, 'Vali')
        self.assertNotContains(resp, 'Toza Mijoz')


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
@mock.patch('core.views.send_tg_msg')
class ClientSideTests(TestCase):
    def setUp(self):
        from .models import AllowedAdmin
        self.owner, self.shop = make_shop()
        AllowedAdmin.objects.create(shop=self.shop, name='Egasi', telegram_id=111111)
        self.vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233', telegram_id=77)
        Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=100000, items='Un', status='confirmed')
        with mock.patch('requests.post'):
            self.debt = Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=50000,
                                            items="Non: 10 x 5 000\nChoy: 1 x 0", status='pending')

    def test_confirm_page_shows_shop_items_and_balance(self, _send):
        resp = self.client.get(reverse('debt_detail', args=[self.debt.uuid]))
        self.assertContains(resp, "Test Do&#x27;kon")
        self.assertContains(resp, 'Non: 10 x 5 000')
        self.assertContains(resp, 'Choy: 1 x 0')
        self.assertContains(resp, "150 000 so&#x27;m (Qarz)")  # tasdiqlangandan keyingi qarz

    def test_confirm_notifies_shop(self, send):
        self.client.post(reverse('debt_detail', args=[self.debt.uuid]),
                         {'action': 'confirm', 'init_data': signed_init(77)})
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.status, 'confirmed')
        send.assert_called_once()
        self.assertEqual(send.call_args[0][0], 111111)
        self.assertIn('tasdiqladi', send.call_args[0][1])

    def test_reject_with_reason_notifies_shop(self, send):
        self.client.post(reverse('debt_detail', args=[self.debt.uuid]),
                         {'action': 'reject', 'reason': 'summa xato', 'init_data': signed_init(77)})
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.status, 'rejected')
        self.assertIn('summa xato', send.call_args[0][1])

    def test_processed_debt_shows_status_instead_of_form(self, _send):
        self.debt.status = 'confirmed'
        self.debt.save()
        resp = self.client.get(reverse('debt_detail', args=[self.debt.uuid]))
        self.assertContains(resp, 'allaqachon tasdiqlangan')
        self.assertNotContains(resp, 'Tasdiqlayman')

    def test_cabinet_shows_pending_and_month_purchases(self, _send):
        Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=-30000, items='tolov',
                            status='confirmed', transaction_type='payment')
        session = self.client.session
        session['client_id'] = self.vali.id
        session.save()
        resp = self.client.get(reverse('client_cabinet'))
        self.assertContains(resp, 'Tasdiqlashingiz kerak (1)')
        self.assertContains(resp, reverse('debt_detail', args=[self.debt.uuid]))
        self.assertContains(resp, "Test Do&#x27;kon")
        self.assertEqual(resp.context['month_debt'], 100000)   # to'lov xaridlarga qo'shilmaydi
        self.assertEqual(resp.context['month_paid'], 30000)


class InitDataTests(TestCase):
    def test_valid_signature(self):
        self.assertEqual(verify_init_data(signed_init(42), TEST_BOT_TOKEN)['id'], 42)

    def test_wrong_token_or_tampered_data_rejected(self):
        self.assertIsNone(verify_init_data(signed_init(42, token='999:other'), TEST_BOT_TOKEN))
        tampered = signed_init(42).replace('42', '43')
        self.assertIsNone(verify_init_data(tampered, TEST_BOT_TOKEN))
        self.assertIsNone(verify_init_data('', TEST_BOT_TOKEN))
        self.assertIsNone(verify_init_data('garbage', TEST_BOT_TOKEN))

    def test_expired_rejected(self):
        self.assertIsNone(verify_init_data(signed_init(42, auth_date=1), TEST_BOT_TOKEN))


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
@mock.patch('requests.post')
class SecurityTests(TestCase):
    def setUp(self):
        self.owner, self.shop = make_shop(tg_id='555555')
        self.vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233', telegram_id=77)

    def auth(self, payload):
        return self.client.post(reverse('telegram_auth'), data=json.dumps(payload), content_type='application/json')

    def test_login_with_plain_telegram_id_is_rejected(self, _post):
        resp = self.auth({'telegram_id': 555555})
        self.assertEqual(resp.status_code, 403)
        self.assertNotIn('_auth_user_id', self.client.session)

    def test_login_with_signed_init_data(self, _post):
        resp = self.auth({'init_data': signed_init(555555)})
        self.assertEqual(resp.json()['redirect_url'], reverse('main_menu'))
        self.assertEqual(self.client.session['_auth_user_id'], str(self.owner.id))

    def test_client_login_with_signed_init_data(self, _post):
        resp = self.auth({'init_data': signed_init(77)})
        self.assertEqual(resp.json()['redirect_url'], reverse('client_cabinet'))
        self.assertEqual(self.client.session['client_id'], self.vali.id)

    def test_login_requires_csrf_token(self, _post):
        from django.test import Client as HttpClient
        strict = HttpClient(enforce_csrf_checks=True)
        resp = strict.post(reverse('telegram_auth'), data=json.dumps({'init_data': signed_init(555555)}),
                           content_type='application/json')
        self.assertEqual(resp.status_code, 403)

    def test_stranger_cannot_confirm_debt(self, _post):
        debt = Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=100, items='x', status='pending')
        url = reverse('debt_detail', args=[debt.uuid])
        self.assertEqual(self.client.post(url, {'action': 'confirm'}).status_code, 403)
        self.assertEqual(self.client.post(url, {'action': 'confirm', 'init_data': signed_init(99)}).status_code, 403)
        debt.refresh_from_db()
        self.assertEqual(debt.status, 'pending')
        with mock.patch('core.views.send_tg_msg'):
            self.client.post(url, {'action': 'confirm', 'init_data': signed_init(77)})
        debt.refresh_from_db()
        self.assertEqual(debt.status, 'confirmed')

    @override_settings(TELEGRAM_WEBHOOK_SECRET='s3cret')
    def test_webhook_requires_secret(self, _post):
        body = json.dumps({'message': {'chat': {'id': 1}, 'text': '/id'}})
        resp = self.client.post(reverse('telegram_webhook'), data=body, content_type='application/json')
        self.assertEqual(resp.status_code, 403)
        with mock.patch('core.views.send_tg_msg'):
            resp = self.client.post(reverse('telegram_webhook'), data=body, content_type='application/json',
                                    HTTP_X_TELEGRAM_BOT_API_SECRET_TOKEN='s3cret')
        self.assertEqual(resp.status_code, 200)


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
@mock.patch('requests.post')
class StaffInviteTests(TestCase):
    def setUp(self):
        self.owner, self.shop = make_shop(tg_id='555555')

    def start(self, token, chat_id=888):
        return self.client.post(reverse('telegram_webhook'), content_type='application/json', data=json.dumps(
            {'message': {'chat': {'id': chat_id}, 'from': {'id': chat_id, 'first_name': 'Ali'}, 'text': f'/start {token}'}}))

    def test_invite_flow(self, _post):
        from .models import StaffInvite
        self.client.force_login(self.owner)
        self.client.post(reverse('manage_admins', args=['invite']), {'name': 'Ali'})
        invite = StaffInvite.objects.get(shop=self.shop)
        resp = self.client.get(reverse('admin_control'))
        self.assertContains(resp, f'staff_{invite.token}')

        with mock.patch('core.views.send_tg_msg'), mock.patch('core.views.send_menu'):
            self.start(f'staff_{invite.token}')
        user = User.objects.get(username='888')
        self.assertEqual(user.profile.shop, self.shop)
        self.assertEqual(user.profile.role, 'worker')
        invite.refresh_from_db()
        self.assertIsNotNone(invite.used_at)

        # Ikkinchi marta ishlatib bo'lmaydi
        with mock.patch('core.views.send_tg_msg') as send, mock.patch('core.views.send_menu'):
            self.start(f'staff_{invite.token}', chat_id=999)
        self.assertFalse(User.objects.filter(username='999').exists())
        self.assertIn('eskirgan', send.call_args[0][1])

    def test_expired_invite(self, _post):
        from .models import StaffInvite
        invite = StaffInvite.objects.create(shop=self.shop, name='Ali')
        StaffInvite.objects.filter(id=invite.id).update(created_at=timezone.now() - timezone.timedelta(days=8))
        with mock.patch('core.views.send_tg_msg'), mock.patch('core.views.send_menu'):
            self.start(f'staff_{invite.token}')
        self.assertFalse(User.objects.filter(username='888').exists())


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
@mock.patch('requests.post')
class MultiShopClientTests(TestCase):
    def setUp(self):
        _, self.shop_a = make_shop(tg_id='111', name='Alfa')
        _, self.shop_b = make_shop(tg_id='222', name='Beta')
        self.a = Client.objects.create(shop=self.shop_a, full_name='Vali', phone='+998901112233', telegram_id=77)
        self.b = Client.objects.create(shop=self.shop_b, full_name='Vali', phone='+998901112233', telegram_id=77)
        self.stranger = Client.objects.create(shop=self.shop_b, full_name='Boshqa', phone='+998900000009', telegram_id=55)

    def test_cabinet_lists_both_shops_and_switches(self, _post):
        self.client.post(reverse('telegram_auth'), data=json.dumps({'init_data': signed_init(77)}),
                         content_type='application/json')
        resp = self.client.get(reverse('client_cabinet'))
        self.assertContains(resp, 'Alfa')
        self.assertContains(resp, 'Beta')
        self.client.get(reverse('client_switch', args=[self.b.id]))
        self.assertEqual(self.client.session['client_id'], self.b.id)
        # Boshqa odamning hisobiga o'tib bo'lmaydi
        self.client.get(reverse('client_switch', args=[self.stranger.id]))
        self.assertEqual(self.client.session['client_id'], self.b.id)

    def test_login_remembers_last_selected_shop(self, _post):
        session = self.client.session
        session['client_id'] = self.b.id
        session.save()
        self.client.post(reverse('telegram_auth'), data=json.dumps({'init_data': signed_init(77)}),
                         content_type='application/json')
        self.assertEqual(self.client.session['client_id'], self.b.id)


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
@mock.patch('requests.post')
class StoreTests(TestCase):
    def setUp(self):
        from .models import AllowedAdmin
        from store.models import Product
        self.owner, self.shop = make_shop(tg_id='555555')
        AllowedAdmin.objects.create(shop=self.shop, name='Egasi', telegram_id=555555)
        _, self.other_shop = make_shop(tg_id='666666', name='Boshqa')
        self.vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233', telegram_id=77)
        self.p1 = Product.objects.create(shop=self.shop, name='Un', price=50000)
        self.foreign = Product.objects.create(shop=self.other_shop, name='Begona', price=1)
        session = self.client.session
        session['client_id'] = self.vali.id
        session.save()

    def test_only_own_shop_products(self, _post):
        resp = self.client.get(reverse('shop_home'))
        self.assertContains(resp, 'Un')
        self.assertNotContains(resp, 'Begona')
        self.assertEqual(self.client.post(reverse('add_to_cart', args=[self.foreign.id])).status_code, 404)

    def test_checkout_and_accept_creates_debt(self, post):
        from store.models import Order
        self.client.post(reverse('add_to_cart', args=[self.p1.id]))
        self.client.post(reverse('add_to_cart', args=[self.p1.id]))
        self.client.post(reverse('checkout'))
        order = Order.objects.get()
        self.assertEqual(order.shop, self.shop)
        self.assertEqual(order.total_price, 100000)
        sent_to = {c.kwargs['json']['chat_id'] for c in post.call_args_list if 'reply_markup' in c.kwargs.get('json', {})}
        self.assertEqual(sent_to, {555555})

        from .views import handle_order_accept
        with mock.patch('core.views.send_tg_msg'), mock.patch('core.views.edit_tg_message'):
            handle_order_accept(999, 1, order.id)          # begona odam
            self.assertFalse(Debt.objects.filter(client=self.vali).exists())
            handle_order_accept(555555, 1, order.id)
            handle_order_accept(555555, 1, order.id)       # ikkinchi bosish
        debts = Debt.objects.filter(client=self.vali)
        self.assertEqual(debts.count(), 1)
        self.assertEqual(debts[0].shop, self.shop)
        self.assertEqual(debts[0].amount_uzs, 100000)

    def test_owner_manages_products_worker_cannot(self, _post):
        from store.models import Product
        self.client.force_login(self.owner)
        self.client.post(reverse('product_add'), {'name': 'Shakar', 'price': '15 000', 'category': 'Oziq', 'is_active': 'on'})
        p = Product.objects.get(name='Shakar')
        self.assertEqual((p.shop, p.price, p.category.name), (self.shop, 15000, 'Oziq'))
        worker = make_worker(self.shop)
        self.client.force_login(worker)
        self.assertRedirects(self.client.get(reverse('manage_products')), reverse('main_menu'))

    def test_pricing_page_is_honest(self, _post):
        self.client.force_login(self.owner)
        resp = self.client.get(reverse('pricing_page'))
        self.assertNotContains(resp, 'SMS')
        self.assertContains(resp, '100 000')


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
class TelegramModuleTests(TestCase):
    def test_every_request_has_timeout(self):
        from . import telegram
        with mock.patch('requests.post') as post:
            post.return_value.json.return_value = {'ok': True}
            telegram.send_message(1, 'salom', background=False)
        self.assertEqual(post.call_args.kwargs['timeout'], telegram.REQUEST_TIMEOUT)
        self.assertEqual(post.call_args.kwargs['json']['text'], 'salom')

    def test_network_error_is_logged_not_raised(self):
        import requests
        from . import telegram
        with mock.patch('requests.post', side_effect=requests.Timeout('sekin')), \
                self.assertLogs('core.telegram', level='WARNING') as logs:
            self.assertIsNone(telegram.send_message(1, 'x', background=False))
        self.assertIn('sekin', logs.output[0])

    def test_background_send_does_not_block(self):
        import threading
        from . import telegram
        release = threading.Event()
        with mock.patch('requests.post', side_effect=lambda *a, **k: release.wait(5)):
            started = timezone.now()
            telegram.send_message(1, 'x', background=True)
            self.assertLess((timezone.now() - started).total_seconds(), 1)
            release.set()

    def test_missing_token_skips_request(self):
        from . import telegram
        with override_settings(BOT_TOKEN=''), mock.patch('requests.post') as post, \
                self.assertLogs('core.telegram', level='WARNING'):
            telegram.send_message(1, 'x', background=False)
        post.assert_not_called()

    def test_broadcast_sends_to_each_linked_client_once(self):
        user, shop = make_shop()
        Client.objects.create(shop=shop, full_name='A', phone='+998900000001', telegram_id=11)
        Client.objects.create(shop=shop, full_name='B', phone='+998900000002', telegram_id=12)
        Client.objects.create(shop=shop, full_name='C', phone='+998900000003')
        self.client.force_login(user)
        with mock.patch('core.telegram.send_message') as send:
            self.client.post(reverse('broadcast'), {'message': 'Aksiya!'})
        self.assertEqual(sorted(c.args[0] for c in send.call_args_list), [11, 12])


class BackupTests(TransactionTestCase):
    """TransactionTestCase: SQLite backup ochiq tranzaksiya ichida ishlamaydi (serverda bu muammo yo'q)."""

    def test_backup_contains_data(self):
        import gzip
        import sqlite3
        import tempfile
        from pathlib import Path
        from django.core.management import call_command

        _, shop = make_shop()
        Client.objects.create(shop=shop, full_name='Vali', phone='+998901112233')
        with tempfile.TemporaryDirectory() as tmp:
            call_command('backup_db', dir=tmp, stdout=mock.Mock())
            [backup] = Path(tmp).glob('db-*.sqlite3.gz')
            restored = Path(tmp) / 'restored.sqlite3'
            with gzip.open(backup) as src:
                restored.write_bytes(src.read())
            con = sqlite3.connect(restored)
            self.assertEqual(con.execute("select full_name from core_client").fetchone()[0], 'Vali')
            con.close()

    def test_rotation_keeps_latest(self):
        import tempfile
        from pathlib import Path
        from core.management.commands.backup_db import Command

        with tempfile.TemporaryDirectory() as tmp:
            for day in range(1, 6):
                (Path(tmp) / f'db-2026010{day}-030000.sqlite3.gz').touch()
            (Path(tmp) / 'boshqa-fayl.txt').touch()
            Command().rotate(Path(tmp), keep=2)
            self.assertEqual(sorted(p.name for p in Path(tmp).iterdir()),
                             ['boshqa-fayl.txt', 'db-20260104-030000.sqlite3.gz', 'db-20260105-030000.sqlite3.gz'])


@mock.patch('requests.post')
class RedesignedPagesTests(TestCase):
    def setUp(self):
        self.owner, self.shop = make_shop()
        self.client.force_login(self.owner)

    def test_settings_updates_rate_and_name(self, _post):
        from .models import Settings
        self.client.post(reverse('settings'), {'action': 'update_rate', 'usd_rate': '12 950'})
        self.assertEqual(Settings.objects.get(shop=self.shop).usd_rate, 12950)
        self.client.post(reverse('settings'), {'action': 'update_rate', 'usd_rate': '5'})  # noto'g'ri
        self.assertEqual(Settings.objects.get(shop=self.shop).usd_rate, 12950)
        self.client.post(reverse('settings'), {'action': 'update_shop', 'shop_name': 'Yangi nom'})
        self.shop.refresh_from_db()
        self.assertEqual(self.shop.name, 'Yangi nom')

    def test_broadcast_shows_recipient_count(self, _post):
        Client.objects.create(shop=self.shop, full_name='A', phone='+998900000001', telegram_id=11)
        Client.objects.create(shop=self.shop, full_name='B', phone='+998900000002')
        resp = self.client.get(reverse('broadcast'))
        self.assertEqual(len(resp.context['clients']), 1)  # faqat botga ulangan

    def test_super_dashboard_and_extend(self, _post):
        self.assertRedirects(self.client.get(reverse('super_dashboard')),
                             '/admin/login/?next=' + reverse('super_dashboard'), fetch_redirect_response=False)
        admin = User.objects.create_superuser('root', password='x')
        self.client.force_login(admin)
        c = Client.objects.create(shop=self.shop, full_name='A', phone='+998900000001')
        Debt.objects.create(shop=self.shop, client=c, amount_uzs=300, items='x', status='confirmed')
        Debt.objects.create(shop=self.shop, client=c, amount_uzs=-100, items='t', status='confirmed',
                            transaction_type='payment')
        resp = self.client.get(reverse('super_dashboard'))
        row = resp.context['rows'][0]
        self.assertEqual((row['shop'].client_count, row['outstanding']), (1, 200))

        self.shop.subscription_ends_at = timezone.now() - timezone.timedelta(days=5)  # tugagan
        self.shop.save()
        self.assertEqual(self.client.get(reverse('extend_subscription', args=[self.shop.id])).status_code, 405)
        self.client.post(reverse('extend_subscription', args=[self.shop.id]))
        self.shop.refresh_from_db()
        self.assertEqual(self.shop.days_left, 29)  # bugundan 30 kun


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
class ReminderTests(TestCase):
    def setUp(self):
        from .models import Settings
        self.owner, self.shop = make_shop()
        self.cfg = Settings.objects.create(shop=self.shop, reminder_enabled=True, reminder_days=7, reminder_min_debt=10000)
        self.debtor = Client.objects.create(shop=self.shop, full_name='Qarzdor', phone='+998900000001', telegram_id=11)
        self.small = Client.objects.create(shop=self.shop, full_name='Kichik', phone='+998900000002', telegram_id=12)
        self.nobot = Client.objects.create(shop=self.shop, full_name='Botsiz', phone='+998900000003')
        with mock.patch('requests.post'):
            for c, amount in [(self.debtor, 50000), (self.small, 5000), (self.nobot, 90000)]:
                Debt.objects.create(shop=self.shop, client=c, amount_uzs=amount, items='x', status='confirmed')

    def run_command(self):
        from django.core.management import call_command
        with mock.patch('core.telegram.send_message') as send, mock.patch('time.sleep'):
            call_command('send_reminders', stdout=mock.Mock())
        return send

    def test_only_due_debtors_with_bot_get_reminder(self):
        send = self.run_command()
        self.assertEqual([c.args[0] for c in send.call_args_list], [11])
        self.assertIn('50 000', send.call_args.args[1])
        # Ertasi kuni - yana yuborilmaydi (7 kun o'tmagan)
        self.assertEqual(self.run_command().call_count, 0)
        # 8 kundan keyin - yana
        Client.objects.filter(id=self.debtor.id).update(last_reminded_at=timezone.now() - timezone.timedelta(days=8))
        self.assertEqual(self.run_command().call_count, 1)

    def test_disabled_or_expired_shop_gets_nothing(self):
        self.cfg.reminder_enabled = False
        self.cfg.save()
        self.assertEqual(self.run_command().call_count, 0)
        self.cfg.reminder_enabled = True
        self.cfg.save()
        self.shop.subscription_ends_at = timezone.now() - timezone.timedelta(days=1)
        self.shop.save()
        self.assertEqual(self.run_command().call_count, 0)

    def test_manual_remind_button_with_cooldown(self):
        self.client.force_login(self.owner)
        with mock.patch('core.telegram.send_message') as send:
            self.client.post(reverse('client_remind', args=[self.debtor.id]))
            self.client.post(reverse('client_remind', args=[self.debtor.id]))  # 1 soat o'tmagan
        self.assertEqual(send.call_count, 1)

    def test_settings_form_saves_reminder_options(self):
        self.client.force_login(self.owner)
        self.client.post(reverse('settings'), {'action': 'update_reminders', 'reminder_days': '3',
                                               'reminder_min_debt': '20 000'})
        self.cfg.refresh_from_db()
        self.assertEqual((self.cfg.reminder_enabled, self.cfg.reminder_days, self.cfg.reminder_min_debt), (False, 3, 20000))


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
@mock.patch('requests.post')
class ExcelExportTests(TestCase):
    def setUp(self):
        self.owner, self.shop = make_shop()
        vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233', telegram_id=5)
        Debt.objects.create(shop=self.shop, client=vali, amount_uzs=300000, items='Un\nShakar', status='confirmed')
        Debt.objects.create(shop=self.shop, client=vali, amount_uzs=-100000, items="To'lov", status='confirmed',
                            transaction_type='payment', payment_method='card')
        self.client.force_login(self.owner)

    def load(self, resp):
        from io import BytesIO
        from openpyxl import load_workbook
        self.assertEqual(resp['Content-Type'], 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        return load_workbook(BytesIO(resp.content))

    def test_clients_export(self, _post):
        ws = self.load(self.client.get(reverse('export_clients'))).active
        self.assertEqual([c.value for c in ws[2]][:3], ['Vali', '+998901112233', 200000])

    def test_month_export(self, _post):
        today = timezone.localdate()
        wb = self.load(self.client.get(reverse('export_month') + f'?date={today:%Y-%m}'))
        ops = list(wb['Operatsiyalar'].iter_rows(min_row=2, values_only=True))
        self.assertEqual([(r[2], r[4]) for r in ops], [('Nasiya', 300000), ("To'lov", 100000)])
        self.assertEqual(list(wb['Mijozlar kesimida'].iter_rows(min_row=2, values_only=True))[0][:3], ('Vali', 300000, 100000))

    def test_send_to_telegram_and_worker_blocked(self, post):
        resp = self.client.get(reverse('export_clients') + '?send=1')
        self.assertRedirects(resp, reverse('dashboard'))
        self.assertTrue(any('sendDocument' in c.args[0] for c in post.call_args_list))
        self.client.force_login(make_worker(self.shop))
        self.assertRedirects(self.client.get(reverse('export_clients')), reverse('main_menu'))


class TargetedBroadcastTests(TestCase):
    def setUp(self):
        self.owner, self.shop = make_shop()
        self.clients = []
        for i, debt in enumerate([500000, 200000, 0, 0]):
            c = Client.objects.create(shop=self.shop, full_name=f'Mijoz{i}', phone=f'+99890000000{i}', telegram_id=100 + i)
            if debt:
                with mock.patch('requests.post'):
                    Debt.objects.create(shop=self.shop, client=c, amount_uzs=debt, items='x', status='confirmed')
            self.clients.append(c)
        _, other = make_shop(tg_id='999999', name='Boshqa')
        self.foreign = Client.objects.create(shop=other, full_name='Begona', phone='+998909999999', telegram_id=555)
        self.client.force_login(self.owner)

    def send(self, **data):
        with mock.patch('core.telegram.send_message') as send:
            self.client.post(reverse('broadcast'), dict({'message': 'Salom {ism}, qarzingiz {qarz}'}, **data))
        return {c.args[0]: c.args[1] for c in send.call_args_list}

    def test_debtors_only(self):
        self.assertEqual(sorted(self.send(mode='debtors')), [100, 101])

    def test_selected_only_and_foreign_ids_ignored(self):
        ids = [str(self.clients[1].id), str(self.clients[3].id), str(self.foreign.id)]
        sent = self.send(mode='selected', client_ids=ids)
        self.assertEqual(sorted(sent), [101, 103])

    def test_personalized_and_escaped(self):
        sent = self.send(mode='selected', client_ids=[str(self.clients[0].id)], message='<b>{ism}</b>: {qarz}')
        self.assertEqual(sent[100], "&lt;b&gt;Mijoz0&lt;/b&gt;: 500 000 so'm")

    def test_nothing_selected(self):
        self.assertEqual(self.send(mode='selected'), {})


@override_settings(BOT_TOKEN=TEST_BOT_TOKEN)
@mock.patch('core.views.send_tg_msg')
class RejectAndResendTests(TestCase):
    """Rad etish sababi do'konga ko'rinadi, do'kon izoh bilan qayta yuboradi."""

    def setUp(self):
        self.owner, self.shop = make_shop()
        self.vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233', telegram_id=77)
        with mock.patch('requests.post'):
            self.debt = Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=50000,
                                            items='Un', status='pending')

    def reject(self, reason='summa xato'):
        self.client.post(reverse('debt_detail', args=[self.debt.uuid]),
                         {'action': 'reject', 'reason': reason, 'init_data': signed_init(77)})
        self.debt.refresh_from_db()

    def test_reason_saved_and_shown_to_shop(self, _send):
        self.reject()
        self.assertEqual(self.debt.reject_reason, 'summa xato')
        self.client.force_login(self.owner)
        detail = self.client.get(reverse('admin_client_detail', args=[self.vali.id]))
        self.assertContains(detail, 'summa xato')
        self.assertContains(detail, 'Izoh bilan qayta yuborish')
        self.assertContains(self.client.get(reverse('main_menu')), 'summa xato')

    def test_resend_rejected_with_note(self, _send):
        self.reject()
        self.client.force_login(self.owner)
        with mock.patch('core.telegram.send_message') as send:
            self.client.post(reverse('manage_debt', args=[self.debt.uuid, 'resend']),
                             {'note': "Summani to'g'riladim"})
        self.debt.refresh_from_db()
        self.assertEqual(self.debt.status, 'pending')
        self.assertEqual(self.debt.shop_note, "Summani to'g'riladim")
        text = send.call_args[0][1]
        self.assertIn('qayta yuborildi', text)
        self.assertIn("Summani to&#x27;g&#x27;riladim", text)  # HTML uchun xavfsiz
        page = self.client.get(reverse('debt_detail', args=[self.debt.uuid]))
        self.assertContains(page, 'summa xato')
        self.assertContains(page, "kon izohi:")

    def test_confirmed_debt_is_not_resent(self, _send):
        self.debt.status = 'confirmed'
        self.debt.save()
        self.client.force_login(self.owner)
        with mock.patch('core.telegram.send_message') as send:
            self.client.post(reverse('manage_debt', args=[self.debt.uuid, 'resend']))
        send.assert_not_called()


@mock.patch('requests.post')
class TabBarAndLiveReloadTests(TestCase):
    def setUp(self):
        self.owner, self.shop = make_shop()
        self.client.force_login(self.owner)
        self.vali = Client.objects.create(shop=self.shop, full_name='Vali', phone='+998901112233', telegram_id=77)

    def test_tabbar_on_all_staff_pages(self, _post):
        for name, tab in [('create_debt', 'Savdo'), ('create_payment', "To&#x27;lov"),
                          ('settings', 'Asosiy'), ('broadcast', 'Asosiy'), ('client_add', 'Mijozlar')]:
            resp = self.client.get(reverse(name))
            self.assertContains(resp, 'class="tabbar"', msg_prefix=name)
        resp = self.client.get(reverse('admin_client_detail', args=[self.vali.id]))
        self.assertContains(resp, 'class="tabbar"')

    def test_no_tabbar_for_client_pages(self, _post):
        self.client.logout()
        session = self.client.session
        session['client_id'] = self.vali.id
        session.save()
        self.assertNotContains(self.client.get(reverse('client_cabinet')), 'class="tabbar"')

    def test_live_version_changes_when_client_confirms(self, _post):
        debt = Debt.objects.create(shop=self.shop, client=self.vali, amount_uzs=1000, items='Un')
        url = reverse('live_version')
        before = self.client.get(url).json()['v']
        self.assertEqual(before, self.client.get(url).json()['v'])
        debt.status = 'confirmed'
        debt.save()
        self.assertNotEqual(before, self.client.get(url).json()['v'])

    def test_live_reload_only_on_list_pages(self, _post):
        self.assertContains(self.client.get(reverse('main_menu')), reverse('live_version'))
        self.assertNotContains(self.client.get(reverse('create_debt')), reverse('live_version'))


@override_settings(BOT_USERNAME='QarzDaptarBot', SUBSCRIPTION_PRICE=100000)
class LandingTests(TestCase):
    def test_landing_points_to_bot_signup(self):
        resp = self.client.get(reverse('landing_page'))
        self.assertContains(resp, 'https://t.me/QarzDaptarBot?start=signup')
        self.assertContains(resp, '100 000')
        self.assertContains(resp, 'property="og:image"')
        self.assertContains(resp, 'lang="uz"')
        self.assertNotContains(resp, 'telegram_id')  # eski forma yo'q

    def test_russian_version_is_remembered(self):
        resp = self.client.get(reverse('landing_page') + '?lang=ru')
        self.assertContains(resp, 'lang="ru"')
        self.assertContains(resp, 'Сантехника')
        self.assertEqual(resp.cookies['site_lang'].value, 'ru')
        self.assertContains(self.client.get(reverse('privacy')), 'Политика конфиденциальности')

    def test_privacy_page(self):
        self.assertContains(self.client.get(reverse('privacy')), 'Maxfiylik siyosati')

    def test_logged_in_shop_owner_skips_landing(self):
        user, _shop = make_shop()
        self.client.force_login(user)
        self.assertRedirects(self.client.get(reverse('landing_page')), reverse('main_menu'))

"""Excel (.xlsx) hisobotlar: mijozlar balansi va oylik operatsiyalar."""
from io import BytesIO

from django.db.models import Max, Q, Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CASH_CLIENT_PHONE, Client, Debt

MONEY = '#,##0'
USD = '#,##0.00'
HEADER_FILL = PatternFill('solid', fgColor='2F7FF0')


def _sheet(wb, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = 'A2'
    return ws


def workbook_bytes(wb):
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _response(wb, filename):
    resp = HttpResponse(workbook_bytes(wb),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


def clients_workbook(shop):
    """Barcha mijozlar va joriy balansi (musbat - qarz, manfiy - haq)."""
    wb = Workbook()
    ws = _sheet(wb, 'Mijozlar', ['F.I.SH', 'Telefon', 'Qarz (so\'m)', 'Qarz ($)', 'Botga ulangan', 'Oxirgi amal'],
                [28, 16, 16, 12, 14, 18], first=True)
    clients = (Client.objects.filter(shop=shop).exclude(phone=CASH_CLIENT_PHONE)
               .annotate(bal_uzs=Sum('debt__amount_uzs', filter=Q(debt__status='confirmed')),
                         bal_usd=Sum('debt__amount_usd', filter=Q(debt__status='confirmed')))
               .order_by('-bal_uzs', 'full_name'))
    last_ops = dict(Debt.objects.filter(shop=shop).values_list('client_id').annotate(last=Max('created_at')))
    total_uzs = total_usd = 0
    for c in clients:
        bal_uzs, bal_usd = c.bal_uzs or 0, c.bal_usd or 0
        total_uzs += bal_uzs
        total_usd += bal_usd
        last = last_ops.get(c.id)
        ws.append([c.full_name, c.phone, bal_uzs, float(bal_usd), 'Ha' if c.telegram_id else "Yo'q",
                   timezone.localtime(last).strftime('%d.%m.%Y %H:%M') if last else ''])
    ws.append([])
    ws.append(['JAMI', '', total_uzs, float(total_usd)])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=4):
        row[0].number_format = MONEY
        row[1].number_format = USD
    return wb


def month_workbook(shop, year, month):
    """Oy bo'yicha barcha operatsiyalar va mijozlar kesimida yig'indi."""
    wb = Workbook()
    ops = (Debt.objects.filter(shop=shop, created_at__year=year, created_at__month=month)
           .select_related('client').order_by('created_at'))
    ws = _sheet(wb, 'Operatsiyalar', ['Sana', 'Mijoz', 'Turi', 'Holati', "So'm", '$', "To'lov usuli", 'Tafsilot'],
                [17, 26, 14, 14, 14, 10, 14, 60], first=True)
    kinds = {('debt', False): 'Nasiya', ('debt', True): 'Naqd savdo', ('payment', False): "To'lov"}
    for op in ops:
        if op.is_cash_sale and op.transaction_type == 'payment':
            continue  # naqd savdoning to'lov jufti - alohida qator shart emas
        ws.append([
            timezone.localtime(op.created_at).strftime('%d.%m.%Y %H:%M'),
            'Naqd savdo' if op.client.phone == CASH_CLIENT_PHONE else op.client.full_name,
            kinds.get((op.transaction_type, op.is_cash_sale), op.transaction_type),
            op.get_status_display(),
            abs(op.amount_uzs), float(abs(op.amount_usd)),
            op.get_payment_method_display() or '',
            op.items.replace('\n', '; ')[:500],
        ])
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
        row[0].number_format = MONEY
        row[1].number_format = USD

    summary = _sheet(wb, 'Mijozlar kesimida', ['Mijoz', "Nasiya (so'm)", "To'lov (so'm)", 'Nasiya ($)', "To'lov ($)"],
                     [28, 16, 16, 12, 12])
    credit = ops.filter(status='confirmed', is_cash_sale=False).exclude(client__phone=CASH_CLIENT_PHONE)
    rows = (credit.values('client__full_name')
            .annotate(d_uzs=Sum('amount_uzs', filter=Q(transaction_type='debt')),
                      p_uzs=Sum('amount_uzs', filter=Q(transaction_type='payment')),
                      d_usd=Sum('amount_usd', filter=Q(transaction_type='debt')),
                      p_usd=Sum('amount_usd', filter=Q(transaction_type='payment')))
            .order_by('client__full_name'))
    for r in rows:
        summary.append([r['client__full_name'], r['d_uzs'] or 0, abs(r['p_uzs'] or 0),
                        float(r['d_usd'] or 0), float(abs(r['p_usd'] or 0))])
    for row in summary.iter_rows(min_row=2, min_col=2, max_col=5):
        row[0].number_format = row[1].number_format = MONEY
        row[2].number_format = row[3].number_format = USD
    return wb


def clients_export(shop):
    return clients_workbook(shop), f"mijozlar-{timezone.localdate():%Y-%m-%d}.xlsx"


def month_export(shop, year, month):
    return month_workbook(shop, year, month), f"hisobot-{year}-{month:02d}.xlsx"


def deliver(request, wb, filename, back_url):
    """
    Brauzerda - fayl yuklab olinadi. Telegram ichida (?send=1) - bot faylni foydalanuvchi chatiga yuboradi,
    chunki Telegram mobil ilovalari Mini App ichidan fayl yuklab olishga ruxsat bermaydi.
    """
    from django.contrib import messages
    from django.shortcuts import redirect
    from . import telegram

    if request.GET.get('send') == '1' and request.user.username.isdigit():
        telegram.send_document(int(request.user.username), filename, workbook_bytes(wb),
                               caption=f"📊 {filename}")
        messages.success(request, "📨 Excel fayl Telegram chatingizga yuborildi.")
        return redirect(back_url)
    return _response(wb, filename)
